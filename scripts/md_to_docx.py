"""Convert DOCUMENTATION.md to DOCUMENTATION.docx with sensible formatting.

Handles:
  - Headings (#, ##, ###, ####)
  - Paragraphs (with inline `code` and **bold** + *italic*)
  - Code blocks (```...```)
  - Pipe tables (| col | col |)
  - Bullet lists (- )
  - Numbered lists (1. )
  - Horizontal rules (---)

Mermaid blocks are rendered as fenced code blocks (Word doesn't render mermaid).
That's expected and documented in the deliverable.
"""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor, Inches, Cm


SRC = Path(__file__).resolve().parent.parent / "DOCUMENTATION.md"
DST = Path(__file__).resolve().parent.parent / "DOCUMENTATION.docx"


# ─────────────────────────── Styles ───────────────────────────

def ensure_code_style(doc):
    styles = doc.styles
    if "OrchestrixCode" in [s.name for s in styles]:
        return styles["OrchestrixCode"]
    style = styles.add_style("OrchestrixCode", WD_STYLE_TYPE.PARAGRAPH)
    font = style.font
    font.name = "Consolas"
    font.size = Pt(9)
    font.color.rgb = RGBColor(0x1F, 0x1F, 0x1F)
    pf = style.paragraph_format
    pf.left_indent = Cm(0.5)
    pf.space_before = Pt(2)
    pf.space_after = Pt(2)
    return style


def ensure_inline_code_style(doc):
    styles = doc.styles
    if "OrchestrixInline" in [s.name for s in styles]:
        return styles["OrchestrixInline"]
    style = styles.add_style("OrchestrixInline", WD_STYLE_TYPE.CHARACTER)
    style.font.name = "Consolas"
    style.font.size = Pt(10)
    style.font.color.rgb = RGBColor(0x9C, 0x27, 0xB0)
    return style


def add_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.find(qn("w:tblPr"))
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    borders = OxmlElement("w:tblBorders")
    for kind in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{kind}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:color"), "999999")
        borders.append(b)
    tblPr.append(borders)


# ─────────────────────────── Inline parsing ───────────────────────────

INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def add_runs(paragraph, text, inline_style):
    """Add text with inline emphasis/code/link parsing."""
    if not text:
        return
    parts = INLINE_RE.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith("`") and part.endswith("`"):
            paragraph.add_run(part[1:-1], style=inline_style)
        elif part.startswith("[") and "](" in part:
            # [text](url) — render as text only (Word hyperlinks need extra plumbing).
            text_part = part[1:part.index("](")]
            paragraph.add_run(text_part).bold = False
        else:
            paragraph.add_run(part)


# ─────────────────────────── Block parsing ───────────────────────────

def is_heading(line):
    m = re.match(r"^(#{1,6})\s+(.*)$", line)
    if not m:
        return None, None
    return len(m.group(1)), m.group(2).strip()


def is_table_row(line):
    return line.strip().startswith("|") and line.strip().endswith("|")


def is_table_separator(line):
    s = line.strip().strip("|").strip()
    return all(re.fullmatch(r":?-+:?", c.strip()) for c in s.split("|") if c.strip())


def parse_table(lines, i):
    """Parse a markdown pipe table starting at index i. Returns (headers, rows, end_index)."""
    header_line = lines[i].strip().strip("|")
    headers = [c.strip() for c in header_line.split("|")]
    # i+1 is the separator line
    rows = []
    j = i + 2
    while j < len(lines) and is_table_row(lines[j]):
        row_line = lines[j].strip().strip("|")
        cells = [c.strip() for c in row_line.split("|")]
        # pad/trim to header size
        cells = (cells + [""] * len(headers))[: len(headers)]
        rows.append(cells)
        j += 1
    return headers, rows, j


# ─────────────────────────── Main converter ───────────────────────────

def convert(md_path: Path, docx_path: Path):
    text = md_path.read_text(encoding="utf-8")
    lines = text.splitlines()

    doc = Document()
    # Tighten default margins.
    for section in doc.sections:
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)

    code_style = ensure_code_style(doc)
    inline_style = ensure_inline_code_style(doc)

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Code block (```...```)
        if stripped.startswith("```"):
            j = i + 1
            buf = []
            while j < len(lines) and not lines[j].strip().startswith("```"):
                buf.append(lines[j])
                j += 1
            for ln in buf:
                p = doc.add_paragraph(ln, style=code_style)
                # No paragraph spacing inside code blocks.
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
            # Add a spacer after the block.
            doc.add_paragraph("")
            i = j + 1
            continue

        # Pipe table
        if (i + 1) < len(lines) and is_table_row(line) and is_table_separator(lines[i + 1]):
            headers, rows, end = parse_table(lines, i)
            t = doc.add_table(rows=1 + len(rows), cols=len(headers))
            t.style = "Light List Accent 1"
            add_table_borders(t)
            for k, h in enumerate(headers):
                cell = t.rows[0].cells[k]
                cell.paragraphs[0].text = ""
                run = cell.paragraphs[0].add_run(h)
                run.bold = True
            for r_i, row in enumerate(rows, start=1):
                for c_i, cell_text in enumerate(row):
                    cell = t.rows[r_i].cells[c_i]
                    cell.paragraphs[0].text = ""
                    add_runs(cell.paragraphs[0], cell_text, inline_style)
            doc.add_paragraph("")
            i = end
            continue

        # Heading
        level, htext = is_heading(line)
        if level is not None:
            level = min(level, 4)
            heading = doc.add_heading(level=level)
            add_runs(heading, htext, inline_style)
            i += 1
            continue

        # Horizontal rule
        if stripped == "---":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run("─" * 60)
            run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
            i += 1
            continue

        # Bullet list
        if stripped.startswith("- ") or stripped.startswith("* "):
            content = stripped[2:]
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, content, inline_style)
            i += 1
            continue

        # Numbered list
        m = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m:
            content = m.group(2)
            p = doc.add_paragraph(style="List Number")
            add_runs(p, content, inline_style)
            i += 1
            continue

        # Blank line
        if not stripped:
            i += 1
            continue

        # Paragraph (may continue across multiple lines until a blank line / new block)
        para_lines = [line]
        j = i + 1
        while j < len(lines):
            nxt = lines[j]
            if not nxt.strip():
                break
            if is_heading(nxt)[0] is not None:
                break
            if nxt.strip().startswith("```"):
                break
            if (j + 1) < len(lines) and is_table_row(nxt) and is_table_separator(lines[j + 1]):
                break
            if nxt.strip().startswith("- ") or nxt.strip().startswith("* "):
                break
            if re.match(r"^\d+\.\s+", nxt.strip()):
                break
            if nxt.strip() == "---":
                break
            para_lines.append(nxt)
            j += 1
        para = " ".join(l.rstrip() for l in para_lines).strip()
        if para:
            p = doc.add_paragraph()
            add_runs(p, para, inline_style)
        i = j

    doc.save(docx_path)
    return docx_path


if __name__ == "__main__":
    out = convert(SRC, DST)
    print(f"wrote {out}")
